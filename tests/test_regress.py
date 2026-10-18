# -*- coding: utf-8 -*-
r"""Регрессионный стенд сам должен быть покрыт тестом.

Стыдно, если страховка, которая ловит регрессии, сломалась тихо: она ведь
единственная, кто смотрит на настоящие выгрузки 1С. Здесь стенд прогоняется
на синтетических фикстурах и проверяется дважды:

* два одинаковых прогона не дают регрессий;
* прогон после подсовывания битого файла — даёт.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from tests import factories
from tools import regress

#: Все двенадцать вариантов выгрузки. Порядок влияет только на порядок отчёта.
BUILDERS = [
    factories.card_upp, factories.card_non_upp,
    factories.posting_upp, factories.posting_non_upp,
    factories.analysis_upp, factories.analysis_non_upp,
    factories.turnover_upp, factories.turnover_non_upp,
    factories.accountosv_upp, factories.accountosv_non_upp,
    factories.generalosv_upp, factories.generalosv_non_upp,
]


@pytest.fixture
def dump_folder(out_dir):
    """Папка со всеми двенадцатью вариантами выгрузок."""
    for number, builder in enumerate(BUILDERS, 1):
        builder(out_dir / f"выгрузка_{number:02d}.xlsx")
    return out_dir


def _run(root: Path, target: Path) -> list:
    """Прогоняет стенд и сохраняет отчёт."""
    records = regress.run(root, "process", only_register=None, limit=0)
    target.write_text(json.dumps(records, ensure_ascii=False, indent=2), encoding="utf-8")
    return records


def test_стенд_обрабатывает_все_двенадцать_вариантов(dump_folder, tmp_path):
    records = _run(dump_folder, tmp_path / "отчёт.json")

    statuses = {record["status"] for record in records}
    assert len(records) == len(BUILDERS)
    assert statuses == {"OK"}, [
        (r["file"], r["status"], r.get("message") or r.get("error"))
        for r in records if r["status"] != "OK"
    ]


def test_стенд_видит_форму_результата(dump_folder, tmp_path):
    """В отчёте обязаны быть число строк и набор колонок.

    Без них стенд не отличает «обработалось» от «обработалось в три раза
    меньше».
    """
    records = _run(dump_folder, tmp_path / "отчёт.json")

    for record in records:
        assert record.get("rows", 0) > 0, record
        assert record.get("columns"), record
        assert "Исх.файл" in record["columns"] or "Имя_файла" in record["columns"], record


def test_одинаковые_прогоны_не_дают_регрессий(dump_folder, tmp_path):
    _run(dump_folder, tmp_path / "до.json")
    _run(dump_folder, tmp_path / "после.json")

    assert regress.assert_no_regression(tmp_path / "до.json",
                                        tmp_path / "после.json") == 0


def test_битый_файл_обнаруживается(dump_folder, tmp_path):
    """Подсунутый мусор обязан поднять код возврата 1."""
    _run(dump_folder, tmp_path / "до.json")

    (dump_folder / "выгрузка_01.xlsx").write_bytes(b"not a zip archive")
    _run(dump_folder, tmp_path / "после.json")

    assert regress.assert_no_regression(tmp_path / "до.json",
                                        tmp_path / "после.json") == 1


def test_пропавшая_строка_обнаруживается(dump_folder, tmp_path):
    """Оборот, ушедший из отчёта, должен быть пойман.

    Это тот случай, ради которого в отчёт добавлены ``rows`` и ``columns``:
    статус остаётся OK, а данные уже не те.
    """
    import openpyxl

    _run(dump_folder, tmp_path / "до.json")

    # Удаляем последнего контрагента из выгрузки ОСВ счёта (выгрузка_09 —
    # девятая в списке BUILDERS). Если из отчёта пропала строка, это должно
    # быть замечено, даже несмотря на статус OK.
    target = dump_folder / "выгрузка_09.xlsx"
    workbook = openpyxl.load_workbook(target)
    sheet = workbook.active
    removed = False
    for row_number in range(sheet.max_row, 0, -1):
        if sheet.cell(row=row_number, column=3).value == "Контрагент 2":
            sheet.delete_rows(row_number)
            removed = True
            break
    workbook.save(target)
    assert removed, "фикстура изменилась: строка «Контрагент 2» не найдена"

    _run(dump_folder, tmp_path / "после.json")

    assert regress.assert_no_regression(tmp_path / "до.json",
                                        tmp_path / "после.json") == 1