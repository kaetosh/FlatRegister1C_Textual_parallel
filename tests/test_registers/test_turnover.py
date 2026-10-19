# -*- coding: utf-8 -*-
"""Обороты счёта.

Здесь проверяется самое интересное место регистра: корреспондирующие счета
живут между денежными блоками, и процессор обязан разложить их по суффиксам
`_до` (дебетовая сторона) и `_ко` (кредитовая). Отсюда же берётся «оборот
после обработки» в сверке.
"""

from __future__ import annotations

import pytest

from flatregister.registers.turnover import Turnover_NonUPPFileProcessor, Turnover_UPPFileProcessor

REGISTER = "turnover"


def _path_for(processor_class, upp_file, non_upp_file):
    return upp_file if processor_class is Turnover_UPPFileProcessor else non_upp_file


@pytest.mark.parametrize("processor_class", [Turnover_UPPFileProcessor,
                                             Turnover_NonUPPFileProcessor])
def test_обороты_счёта_разворачивают_иерархию(processor_class, turnover_upp_file,
                                            turnover_non_upp_file):
    path = _path_for(processor_class, turnover_upp_file, turnover_non_upp_file)

    result, check = processor_class().process_file(path)

    assert not result.empty
    assert "Субконто" in result.columns
    # Родительский счёт 60.01 и строка «Итого» отбрасываются: их обороты уже
    # учтены в листьях.
    assert set(result["Субконто"]) == {"Поставщик А", "Поставщик Б"}
    assert set(result["Исх.файл"]) == {path.name}


@pytest.mark.parametrize("processor_class", [Turnover_UPPFileProcessor,
                                             Turnover_NonUPPFileProcessor])
def test_обороты_разносят_коррсчета_по_сторонам(processor_class, turnover_upp_file,
                                                turnover_non_upp_file):
    """Колонки корр. счетов получают суффиксы `_до` и `_ко` по стороне баланса."""
    path = _path_for(processor_class, turnover_upp_file, turnover_non_upp_file)

    result, _ = processor_class().process_file(path)

    debit_side = {column for column in result.columns if column.endswith("_до")}
    credit_side = {column for column in result.columns if column.endswith("_ко")}
    assert debit_side, "не найдено ни одной дебетовой колонки корр. счетов"
    assert credit_side, "не найдено ни одной кредитной колонки корр. счетов"
    assert debit_side.isdisjoint(credit_side)


@pytest.mark.parametrize("processor_class", [Turnover_UPPFileProcessor,
                                             Turnover_NonUPPFileProcessor])
def test_обороты_сверка_нулевая(processor_class, turnover_upp_file,
                              turnover_non_upp_file):
    """Обороты не теряются и не задваиваются.

    «Оборот после» считается как сумма колонок `_до` минус сумма `_ко` по
    оставшимся строкам. Если фикстура собрана верно, разница обязана быть
    нулём — иначе расхождение означало бы дыру в алгоритме.
    """
    path = _path_for(processor_class, turnover_upp_file, turnover_non_upp_file)

    result, check = processor_class().process_file(path)

    assert not check.empty
    row = check.iloc[0]
    detail = f"сверка: {row.to_dict()}\nкадр:\n{result.to_string()}"
    assert row["Сальдо_начало_разница"] == 0, detail
    assert row["Оборот_разница"] == 0, detail
    assert row["Сальдо_конец_разница"] == 0, detail
    assert row["Исх.файл"] == path.name


@pytest.mark.parametrize("processor_class", [Turnover_UPPFileProcessor,
                                             Turnover_NonUPPFileProcessor])
def test_обороты_денежные_блоки_переименованы(processor_class, turnover_upp_file,
                                             turnover_non_upp_file):
    """Шесть денежных блоков приводятся к общим именам независимо от конфигурации."""
    path = _path_for(processor_class, turnover_upp_file, turnover_non_upp_file)

    result, _ = processor_class().process_file(path)

    for column in ("Дебет_начало", "Кредит_начало", "Дебет_оборот",
                   "Кредит_оборот", "Дебет_конец", "Кредит_конец"):
        assert column in result.columns, sorted(result.columns)
    # Служебная колонка предобработки наружу не уходит
    assert "Уровень" not in result.columns
    assert "Курсив" not in result.columns


def test_обороты_без_строки_итого_даёт_понятную_ошибку(out_dir):
    """Нет «Итого» — сверять нечего, но падать внутри pandas нельзя."""
    import openpyxl

    from tests import factories

    path = factories.turnover_upp(out_dir / "обороты без итога.xlsx")
    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    for row_number in range(sheet.max_row, 0, -1):
        if sheet.cell(row=row_number, column=3).value == "Итого":
            sheet.delete_rows(row_number)
            break
    workbook.save(path)

    with pytest.raises(Exception) as caught:
        Turnover_UPPFileProcessor().process_file(path)

    error = caught.value
    assert "Итого" in error.user_message, error.user_message
    assert error.hint, "у структурной ошибки обязана быть подсказка"