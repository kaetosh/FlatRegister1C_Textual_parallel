# -*- coding: utf-8 -*-
"""ОСВ счёта: главное здесь — сверка оборотов до и после обработки."""

from __future__ import annotations

import pytest

from flatregister.registers.account_osv import AccountOSV_NonUPPFileProcessor, AccountOSV_UPPFileProcessor

REGISTER = "accountosv"


def _path_for(processor_class, upp_file, non_upp_file):
    return upp_file if processor_class is AccountOSV_UPPFileProcessor else non_upp_file


@pytest.mark.parametrize("processor_class", [AccountOSV_UPPFileProcessor,
                                             AccountOSV_NonUPPFileProcessor])
def test_осв_счёта_разворачивает_иерархию(processor_class, accountosv_upp_file,
                                       accountosv_non_upp_file):
    path = _path_for(processor_class, accountosv_upp_file, accountosv_non_upp_file)

    result, check = processor_class().process_file(path)

    assert not result.empty
    assert "Субконто" in result.columns
    # Родительская строка 60.01 отбрасывается: её обороты уже учтены в
    # строках контрагентов, иначе сумма задвоится.
    assert set(result["Субконто"]) == {"Контрагент 1", "Контрагент 2"}
    assert "Исх.файл" in result.columns
    assert set(result["Исх.файл"]) == {path.name}
    assert not any(str(value).startswith("Итого") for value in result["Субконто"])


@pytest.mark.parametrize("processor_class", [AccountOSV_UPPFileProcessor,
                                             AccountOSV_NonUPPFileProcessor])
def test_осв_счёта_сверка_нулевая(processor_class, accountosv_upp_file,
                                accountosv_non_upp_file):
    """Обороты не должны ни теряться, ни задваиваться.

    Расхождение «до минус после» и есть тот сигнал, ради которого существует
    лист проверки. На корректной фикстуре оно обязано быть нулём.
    """
    path = _path_for(processor_class, accountosv_upp_file, accountosv_non_upp_file)

    _, check = processor_class().process_file(path)

    assert not check.empty
    row = check.iloc[0]
    assert row["Сальдо_начало_разница"] == 0, row.to_dict()
    assert row["Оборот_разница"] == 0, row.to_dict()
    assert row["Сальдо_конец_разница"] == 0, row.to_dict()
    assert row["Исх.файл"] == path.name


def test_осв_счёта_упп_берёт_номер_счёта_из_заголовка():
    """Номер счёта вытаскивается из строки «... ведомость по счету 60.01».

    Он нужен для выгрузок, где счёт в разбивке не написан: тогда единственный
    источник номера — заголовок отчёта.
    """
    import pandas as pd

    from flatregister.registers.account_osv import find_account_from_text

    frame = pd.DataFrame({"a": ["Оборотно-сальдовая ведомость по счету 60.01"], "b": [1]})
    assert find_account_from_text(frame) == "60.01"

    assert find_account_from_text(pd.DataFrame({"a": ["ничего похожего"], "b": [1]})) is None


def test_осв_счёта_без_строки_итого_даёт_понятную_ошибку(out_dir):
    """Нет «Итого» — сверять нечего, но падать внутри pandas нельзя."""
    from tests import factories

    path = factories.accountosv_upp(out_dir / "без итога.xlsx", rows=[
        ("60.01", 1, 1_000, 2_000, 5_000, 4_000),
    ])
    # Убираем строку «Итого» из готового файла: фабрика её всегда добавляет,
    # поэтому здесь собираем вручную тем же каркасом.
    import openpyxl

    workbook = openpyxl.load_workbook(path)
    sheet = workbook.active
    for row_number in range(sheet.max_row, 0, -1):
        if sheet.cell(row=row_number, column=3).value == "Итого":
            sheet.delete_rows(row_number)
            break
    workbook.save(path)

    with pytest.raises(Exception) as caught:
        AccountOSV_UPPFileProcessor().process_file(path)

    error = caught.value
    assert "Итого" in error.user_message, error.user_message
    assert error.hint, "у структурной ошибки обязана быть подсказка"


def test_осв_счёта_колонки_денежные_переименованы(accountosv_upp_processed):
    """Денежные блоки из двухколоночных групп разворачиваются в шесть колонок."""
    result, _ = accountosv_upp_processed
    for column in ("Дебет_начало", "Кредит_начало", "Дебет_оборот",
                   "Кредит_оборот", "Дебет_конец", "Кредит_конец"):
        assert column in result.columns, sorted(result.columns)
    # Служебная колонка предобработки наружу не уходит
    assert "Уровень" not in result.columns